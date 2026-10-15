from pathlib import Path

from openpyxl import load_workbook

from ExcelWorkbook import ExcelWorkbook


class ExcelSheetReader:
    """Czyta pliki .xlsx (openpyxl) do rekordu ExcelWorkbook."""

    def reader(self) -> "ExcelSheetReader":
        return self

    def read(self, path: Path) -> ExcelWorkbook:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = {sheet.title: [tuple(row) for row in sheet.iter_rows(values_only=True)]
                      for sheet in workbook.worksheets}
        finally:
            workbook.close()
        return ExcelWorkbook(name=Path(path).name, sheets=sheets)
